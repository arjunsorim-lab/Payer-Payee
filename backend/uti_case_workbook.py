"""Reader for the patient-specific UTI billed-cost calculation workbook."""

from pathlib import Path
import os
import re

from openpyxl import load_workbook


DEFAULT_UTI_WORKBOOK = Path("/Users/user/Downloads/Patient_fe621c76_UTI_Culture_Savings_Calc.xlsx")
DEFAULT_SUPPORTING_DATASETS = {
    "synthetic_demo": Path("/Users/user/Downloads/UTI_Value_Based_Case_SYNTHETIC_All_Requirements_Yes.xlsx"),
    "requirements_pack": Path("/Users/user/Downloads/UTI_Value_Based_Case_Complete_Requirements_Pack.xlsx"),
}


def _amount(value):
    if isinstance(value, (int, float)):
        return round(float(value), 2)
    match = re.search(r"-?[0-9][0-9,]*(?:\.[0-9]+)?", str(value or ""))
    return round(float(match.group(0).replace(",", "")), 2) if match else None


def load_uti_case_workbook(path=None):
    configured = path or os.getenv("UTI_CASE_WORKBOOK_PATH", "").strip()
    workbook_path = Path(configured).expanduser() if configured else DEFAULT_UTI_WORKBOOK
    if not workbook_path.is_file():
        return {"available": False, "reason": "The patient UTI calculation workbook was not found.", "source_path": str(workbook_path)}

    sheet = load_workbook(workbook_path, data_only=True, read_only=True).active
    rows = {str(row[0]).strip(): row for row in sheet.iter_rows(values_only=True) if row and row[0]}
    title = str(sheet.cell(1, 1).value or "")
    patient_match = re.search(r"Patient (.+?) \(ID ([^)]+)\)", title)

    def row_amount(label):
        row = rows.get(label)
        if not row:
            return None
        return _amount(row[1])

    earlier_actual = row_amount("Episode 1 Subtotal (Actual, no culture)")
    add_on = row_amount("Culture Add-on Subtotal")
    proposed = row_amount("Episode 1 Subtotal + Urine Culture & Specimen Collection")
    later_actual = row_amount("Episode 2 Subtotal (Total Actual)")
    actual_total = row_amount("Actual Cost")
    savings = row_amount("Savings")
    return {
        "available": True,
        "source": {"workbook_name": workbook_path.name, "workbook_path": str(workbook_path), "sheet": sheet.title},
        "patient": {"name": patient_match.group(1) if patient_match else "Patient-specific UTI case", "id": patient_match.group(2) if patient_match else None},
        "calculation": {
            "earlier_episode_actual_billed": earlier_actual,
            "later_episode_actual_billed": later_actual,
            "culture_and_specimen_add_on_billed": add_on,
            "actual_two_episode_billed": actual_total,
            "proposed_earlier_episode_with_add_on_billed": proposed,
            "potential_billed_difference": savings,
            "formula": "Actual Episode 1 + Episode 2 billed cost − Episode 1 with culture/specimen add-on billed cost",
        },
        "method": "Values are read from the workbook's Billed column and Savings row; this app does not recalculate or replace the workbook result.",
        "quality_of_care": {
            "goal": "Identify and manage the urinary infection earlier so the patient has better diagnostic evidence and a lower risk of recurrence or escalation.",
            "recommended_intervention": "At the earlier UTI episode, review whether the urine culture and specimen-collection bundle should have been performed, then use the result to guide targeted treatment and follow-up.",
            "patient_benefit": "A culture can provide organism and susceptibility evidence that supports more targeted treatment and follow-up instead of waiting for a later recurrence visit.",
            "cost_connection": "The workbook models how earlier diagnostic work could be associated with a lower total billed amount if a later recurrence episode did not occur. It does not prove that the intervention prevented the recurrence or that the amount is confirmed savings.",
            "application_scope": "Apply this quality-of-care review to each member's own UTI episodes and billed lines. Do not copy the attached patient's dollar amounts to other members.",
        },
    }


def dataset_registry():
    """Describe the three uploaded UTI workbooks without mixing their roles."""
    calculation_path = Path(os.getenv("UTI_CASE_WORKBOOK_PATH", "").strip()).expanduser() if os.getenv("UTI_CASE_WORKBOOK_PATH", "").strip() else DEFAULT_UTI_WORKBOOK
    entries = [
        ("patient_calculation", calculation_path, "authoritative billed calculation"),
        ("synthetic_demo", DEFAULT_SUPPORTING_DATASETS["synthetic_demo"], "synthetic demonstration evidence only"),
        ("requirements_pack", DEFAULT_SUPPORTING_DATASETS["requirements_pack"], "requirements and evidence-gap documentation"),
    ]
    return [
        {"dataset_id": dataset_id, "workbook_name": path.name, "role": role, "available": path.is_file()}
        for dataset_id, path, role in entries
    ]
